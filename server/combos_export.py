#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""导出命轮组合明细 Excel。

用法：python3 combos_export.py 输出路径.xlsx
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent


def main() -> None:
    out_path = Path(sys.argv[1])
    proc = subprocess.run(
        [sys.executable, str(HERE / "combos_data.py")],
        capture_output=True, text=True, timeout=60,
    )
    if proc.returncode != 0:
        sys.stderr.write(proc.stderr)
        raise SystemExit(1)
    data = json.loads(proc.stdout)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "组合明细"

    headers = [
        "元素", "轮盘", "组合名", "表定星级", "成员", "成员池",
        "基础效果（每升1星）", "进阶效果（每次进阶）",
        "命轮碎片·单轮/人", "命轮碎片·单轮合计", "命轮碎片·拉满5轮合计",
        "进阶角色碎片/人/次（限定SSR·UR）", "数据来源",
    ]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FCD34D")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")

    for c in data["combos"]:
        members = "、".join(m["name"] for m in c["members"])
        pools = "、".join(f"{m['name']}({m['pool']})" for m in c["members"])
        adv_costs = "、".join(f"{m['name']} {m['advance_cost']}" for m in c["members"] if m["advance_cost"])
        ws.append([
            c["element"], c["wheel"], c["name"], c["stars"], members, pools,
            "；".join(c["effects"]), "；".join(c["advance_effects"]),
            c["wheel_per_member_round"], c["wheel_round_total"], c["wheel_full_total"],
            adv_costs or "—（无限定/UR成员，默认无限）",
            c["source"],
        ])

    widths = [6, 8, 18, 9, 30, 40, 40, 40, 12, 14, 16, 30, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    ws2 = wb.create_sheet("说明")
    notes = [
        ["龙族：卡塞尔之门 · 命轮组合明细"],
        [""],
        ["数据支持：君度"],
        ["技术支持：Roy、梧桐落"],
        [""],
        ["工作簿版本指纹（SHA-256 前12位）", data["workbook_sha256"]],
        ["收录组合总数", data["combo_total"]],
        ["分元素", "；".join(f"{k} {v}" for k, v in data["combo_counts"].items())],
        [""],
        ["口径"],
        ["品阶", "白→绿→蓝→紫→橙→彩，彩0星为终点，共 5 次进阶"],
        ["基础效果", "每升 1 星触发一次，升到表定星级即按 星级×表值 累计"],
        ["进阶效果", "每完成一次真实进阶计一次（当前仅精神元素页提供）"],
        ["命轮碎片", "单轮/人 = 表定星级；拉满 5 轮/人 = 表定星级×5"],
        ["角色碎片", "进阶用：绘梨衣 15/次，其余限定 SSR 与 UR 30/次；常驻 SSR/SR/R 默认视为无限"],
        ["缺失字段", "基础页未收录的属性保持未知，不推断"],
    ]
    for row in notes:
        ws2.append(row)
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 80

    wb.save(out_path)
    print(str(out_path))


if __name__ == "__main__":
    main()
